import docx, os
src = r"C:\Users\Heart\Desktop\艺谋学长课程配套文件\艺谋学长的爆款作品物料\《你的两千年修行之路》04集提示词\第四集第2至4段_14个独立生成单元_完整提示词.docx"
d = docx.Document(src)
lines = [p.text.rstrip() for p in d.paragraphs if p.text.strip()]
out = r"C:\\Users\\Heart\\PycharmProjects\\vomin_studio\\workspace\\养妹失策，社恐妹妹占有欲爆棚\_audit\ref_14units.txt"
open(out, "w", encoding="utf-8").write("\n".join(lines))
print("paras:", len(lines), "chars:", sum(len(x) for x in lines), "tables:", len(d.tables))
