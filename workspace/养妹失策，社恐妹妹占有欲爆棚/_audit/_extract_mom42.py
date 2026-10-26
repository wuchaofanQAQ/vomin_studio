import docx, sys, os
src = r"C:\Users\Heart\Desktop\艺谋学长课程配套文件\剧本转提示词专用文件夹：2026年10月更新\剧本转提示词4.2.docx"
out = r"C:\\Users\\Heart\\PycharmProjects\\vomin_studio\\workspace\\养妹失策，社恐妹妹占有欲爆棚\\_audit\mom42_raw.txt"
d = docx.Document(src)
lines = []
for p in d.paragraphs:
    t = p.text.rstrip()
    if t:
        lines.append(t)
text = "\n".join(lines)
open(out, "w", encoding="utf-8").write(text)
print("paragraphs:", len(d.paragraphs), "chars:", len(text), "tables:", len(d.tables))
for i, tb in enumerate(d.tables[:3]):
    print("table", i, "rows", len(tb.rows), "cols", len(tb.columns))
