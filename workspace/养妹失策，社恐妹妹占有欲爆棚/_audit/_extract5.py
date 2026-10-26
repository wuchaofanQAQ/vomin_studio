import docx, os
src = r"C:\Users\Heart\Desktop\艺谋学长课程配套文件\艺谋学长的爆款作品物料\《你的两千年修行之路》05集提示词+武戏改良版\你的两千年修行之路 第五集 Seedance2.5导演级视频提示词.docx"
d = docx.Document(src)
lines = [p.text.rstrip() for p in d.paragraphs if p.text.strip()]
out = r"C:\\Users\\Heart\\PycharmProjects\\vomin_studio\\workspace\\养妹失策，社恐妹妹占有欲爆棚\_audit\ref_ep5.txt"
open(out, "w", encoding="utf-8").write("\n".join(lines))
print("paras:", len(lines), "chars:", sum(len(x) for x in lines), "tables:", len(d.tables))
